"""Detect page/slide count from uploaded documents."""

import os
import tempfile

from .utils import ALLOWED_UPLOAD_EXTENSIONS


def detect_pages(uploaded_file):
    """
    Return dict: pages (int), method (str), confidence (str: high|medium|low).
    """
    ext = os.path.splitext(uploaded_file.name)[1].lower()
    if ext not in ALLOWED_UPLOAD_EXTENSIONS:
        return {'pages': 1, 'method': 'default', 'confidence': 'low', 'error': 'Unsupported file type'}

    if ext in ('.jpg', '.jpeg', '.png'):
        return {'pages': 1, 'method': 'image', 'confidence': 'high'}

    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
        try:
            for chunk in uploaded_file.chunks():
                tmp.write(chunk)
            tmp.flush()
            path = tmp.name
            uploaded_file.seek(0)

            if ext == '.pdf':
                return _detect_pdf(path)
            if ext == '.docx':
                return _detect_docx(path)
            if ext in ('.ppt', '.pptx'):
                return _detect_pptx(path)
            if ext == '.doc':
                return {'pages': 1, 'method': 'legacy_doc', 'confidence': 'low',
                        'error': 'Legacy .doc — enter page count manually'}
            return {'pages': 1, 'method': 'default', 'confidence': 'low'}
        finally:
            try:
                os.unlink(tmp.name)
            except OSError:
                pass


def _detect_pdf(path):
    try:
        from pypdf import PdfReader
        reader = PdfReader(path)
        pages = len(reader.pages)
        return {'pages': max(1, pages), 'method': 'pdf', 'confidence': 'high'}
    except Exception as exc:
        return {'pages': 1, 'method': 'pdf', 'confidence': 'low', 'error': str(exc)[:120]}


def _detect_docx(path):
    try:
        from docx import Document
        doc = Document(path)
        words = sum(len(p.text.split()) for p in doc.paragraphs)
        pages = max(1, round(words / 250)) if words else 1
        return {'pages': pages, 'method': 'docx_estimate', 'confidence': 'medium',
                'word_count': words}
    except Exception as exc:
        return {'pages': 1, 'method': 'docx', 'confidence': 'low', 'error': str(exc)[:120]}


def _detect_pptx(path):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        slides = len(prs.slides)
        return {'pages': max(1, slides), 'method': 'pptx_slides', 'confidence': 'high'}
    except Exception as exc:
        return {'pages': 1, 'method': 'pptx', 'confidence': 'low', 'error': str(exc)[:120]}
